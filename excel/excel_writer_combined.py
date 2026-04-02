import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, NamedStyle
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from decimal import Decimal
import os



def compute_m0_total_from_env(expanded_env: dict) -> float:
    total = 0.0
    for env in ("Dev", "QA", "Prod"):
        v = expanded_env.get(env, {}).get("M1", 0)
        if isinstance(v, (int, float)):
            total += v
    return round(total, 2)


def scale_cost_components_to_total(cost_components: list, target_total: float):
    if not target_total or target_total <= 0:
        return cost_components

    values = []
    for c in cost_components:
        try:
            values.append(float(c.get("cost_usd", 0)))
        except:
            values.append(0.0)

    current_total = sum(values)
    if current_total == 0:
        return cost_components

    scale = target_total / current_total

    scaled = []
    for comp in cost_components:
        new_comp = comp.copy()
        try:
            val = float(comp.get("cost_usd", 0))
            new_comp["cost_usd"] = round(val * scale, 2)
        except:
            pass
        scaled.append(new_comp)

    return scaled


def scale_env_costs_to_budget(monthly_env: dict, budget: float):
    if not budget or budget <= 0:
        return monthly_env

    def to_number(v):
        try:
            return float(v)
        except:
            return None

    current_total = 0.0
    for env_data in monthly_env.values():
        for i in range(1, 13):
            v = to_number(env_data.get(f"M{i}", 0))
            if v is not None:
                current_total += v

    if current_total == 0:
        return monthly_env

    scale = budget / current_total

    scaled = {}
    for env, env_data in monthly_env.items():
        new_env = {}
        env_sum = 0.0

        for k, v in env_data.items():
            num = to_number(v)
            if k.startswith("M") and num is not None:
                new_val = round(num * scale, 2)
                new_env[k] = new_val
                env_sum += new_val
            else:
                new_env[k] = v

        new_env["Total"] = round(env_sum, 2)
        scaled[env] = new_env

    return scaled


# Excel helpers

def _try_number(v):
    if v is None:
        return ""
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip().replace(",", "")
    if s == "":
        return ""
    try:
        return float(s) if "." in s else int(s)
    except:
        try:
            return float(Decimal(s))
        except:
            return v


def _normalize_string(v):
    if v is None:
        return ""
    if isinstance(v, list):
        return ", ".join(str(x) for x in v)
    if isinstance(v, dict):
        return "; ".join(f"{k}: {v}" for k, v in v.items())
    return str(v)


def auto_fit_columns(ws, extra_padding=4):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value)) if cell.value else 0)
            except:
                pass
        ws.column_dimensions[col_letter].width = max_len + extra_padding


HEADER_FILL = PatternFill(start_color="ADD8E6", end_color="ADD8E6", fill_type="solid")
HEADER_FONT = Font(bold=True)
TITLE_FONT = Font(size=14, bold=True)

CURRENCY = NamedStyle(name="currency_style")
CURRENCY.number_format = '"$"#,##0.00'


# Architecture sheet
def write_architecture_diagram_sheet(wb, image_path, use_case_name):
    ws = wb.create_sheet(f"Architecture_{use_case_name}")
    ws["A1"] = "Architecture Diagram"
    ws["A1"].font = TITLE_FONT

    if not image_path or not os.path.exists(image_path):
        ws["A3"] = "Architecture diagram not available."
        return ws

    img = XLImage(image_path)
    img.width = min(img.width, 900)
    img.height = min(img.height, 600)

    ws.add_image(img, "A3")
    ws.column_dimensions["A"].width = 120
    ws.row_dimensions[3].height = 400
    return ws


# Baseline sheet
def write_combined_sheet(wb, baseline_list, cost_components, pipelines):

    ws = wb.create_sheet("Baseline_cost_assumption")
    row = 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.cell(row=row, column=1).value = "CLOUD COST ESTIMATION OVERVIEW"
    ws.cell(row=row, column=1).font = TITLE_FONT
    ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")
    row += 2

    # BASELINE SUMMARY
    ws.cell(row=row, column=1).value = "Baseline Summary"
    ws.cell(row=row, column=1).font = HEADER_FONT
    row += 1

    headers = ["Parameter", "Usecase Details", "Source of Assumption", "Notes"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    row += 1

    for entry in baseline_list:
        ws.append([
            entry.get("parameter", ""),
            _normalize_string(entry.get("usecase_details", "")),
            _normalize_string(entry.get("source_of_assumption", "")),
            _normalize_string(entry.get("notes", ""))
        ])
        row += 1

    row += 2

    # COST COMPONENTS
    ws.cell(row=row, column=1).value = "Detailed Cost Components"
    ws.cell(row=row, column=1).font = HEADER_FONT
    row += 1

    headers = ["Cost Component", "Calculation Logic", "Cost ($)", "Source", "Remarks"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    row += 1

    for comp in cost_components:
        cost_val = _try_number(comp.get("cost_usd", ""))
        ws.append([
            comp.get("component", ""),
            comp.get("calculation_logic", ""),
            cost_val,
            comp.get("source", ""),
            comp.get("remarks", "")
        ])
        cell = ws.cell(row=row, column=3)
        if isinstance(cost_val, (float, int)):
            cell.number_format = CURRENCY.number_format
        row += 1

    row += 2

    # PIPELINE GROUPS  
    ws.cell(row=row, column=1).value = "Pipeline Groups"
    ws.cell(row=row, column=1).font = HEADER_FONT
    row += 1

    headers = [
        "Pipeline Group", "Data Sources Included", "Refresh Frequency",
        "Runs/Month", "Avg Hours/Run", "Total Hours/Month"
    ]

    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    row += 1

    for p in pipelines:
        ws.append([
            p.get("pipeline_name", ""),
            _normalize_string(p.get("data_sources_included", "")),
            p.get("refresh_frequency", ""),
            _try_number(p.get("runs_per_month", "")),
            _try_number(p.get("avg_hours_per_run", "")),
            _try_number(p.get("total_hours_per_month", ""))
        ])
        row += 1

    auto_fit_columns(ws)
    return ws


# Yearly cost sheet

def write_monthly_environment_sheet(wb, monthly_env, markets=None, global_consumption_multiplier=1.0, budget=None):
    ws = wb.create_sheet("Yearly_Cost")

    market_timeline = [{"market": "M0", "start_month": 1}]
    if markets:
        market_timeline.extend(markets)

    active_markets_per_month = {}
    multiplier_per_month = {}

    for m in range(1, 13):
        active_markets_per_month[m] = "+".join(mk["market"] for mk in market_timeline if mk["start_month"] <= m)
        multiplier_per_month[m] = sum(global_consumption_multiplier for mk in market_timeline if mk["start_month"] <= m)

    expanded_env = {}
    for env in ("Dev", "QA", "Prod"):
        expanded_env[env] = {}
        for i in range(1, 13):
            base = _try_number(monthly_env.get(env, {}).get(f"M{i}", 0))
            expanded_env[env][f"M{i}"] = round(base * multiplier_per_month[i], 2)

    if budget:
        expanded_env = scale_env_costs_to_budget(expanded_env, budget)

    ws.append(["", "Months"] + [f"m{i}" for i in range(1, 13)] + ["Total"])
    ws.append(["", "Markets"] + [active_markets_per_month[i] for i in range(1, 13)] + [""])

    for r in (1, 2):
        for cell in ws[r]:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL

    for env in ("Dev", "QA", "Prod"):
        months = [expanded_env[env].get(f"M{i}", 0) for i in range(1, 13)]
        total = round(sum(months), 2)
        ws.append(["", env] + months + [total])

    ws["A3"] = "Environment"
    ws["A3"].font = HEADER_FONT
    ws["A3"].fill = HEADER_FILL

    auto_fit_columns(ws)
    return expanded_env


# Main entry
def generate_cost_excel_combined(
    json_output,
    file_path,
    client_name,
    use_case_name,
    architecture_image_path=None,
    markets=None,
    global_consumption_multiplier=1.0,
    budget=None
):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    baseline = json_output.get("baseline_summary", [])
    cost_components = json_output.get("detailed_cost_components", [])
    pipelines = json_output.get("pipeline_groups", [])
    monthly_env = json_output.get("monthly_environment_costs", {})

    if architecture_image_path:
        write_architecture_diagram_sheet(wb, architecture_image_path, use_case_name)

    expanded_env = write_monthly_environment_sheet(
        wb,
        monthly_env,
        markets,
        global_consumption_multiplier,
        budget
    )

    m0_total = compute_m0_total_from_env(expanded_env)

    cost_components = scale_cost_components_to_total(cost_components, m0_total)

    write_combined_sheet(wb, baseline, cost_components, pipelines)

    parent_dir = os.path.dirname(file_path)
    if parent_dir:
        os.makedirs(parent_dir, exist_ok=True)

    wb.save(file_path)
    return file_path
